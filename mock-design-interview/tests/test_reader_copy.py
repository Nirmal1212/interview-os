import filecmp
import os
import sys
import tempfile
import unittest

HERE = os.path.dirname(__file__)
PLUGIN_CORE = os.path.abspath(os.path.join(HERE, "..", "skills", "mock-design-interview", "core"))
RESUME_READER = os.path.abspath(
    os.path.join(HERE, "..", "..", "resume-optimization", "skills", "ats-resume-optimizer", "core", "docx_io.py")
)
sys.path.insert(0, PLUGIN_CORE)

from docx import Document  # noqa: E402

import docx_io  # noqa: E402


class ReaderCopyTests(unittest.TestCase):
    def test_reader_is_byte_identical_to_resume_plugin_copy(self):
        # Installed plugins are isolated, so the reader is copied, not imported.
        # This test fails if the two copies drift apart.
        self.assertTrue(
            filecmp.cmp(RESUME_READER, os.path.join(PLUGIN_CORE, "docx_io.py"), shallow=False),
            "docx_io.py differs from the resume plugin copy; re-copy it byte-for-byte",
        )

    def test_reads_headings_from_a_docx_resume(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "resume.docx")
            document = Document()
            document.add_heading("Experience", level=2)
            document.add_paragraph("Built batch and streaming pipelines for four product teams.")
            document.add_heading("Skills", level=2)
            document.add_paragraph("Python, SQL, Spark, Airflow")
            document.save(path)

            blocks = docx_io.read_resume(path)

        headings = [b.text for b in blocks if b.kind == "heading"]
        self.assertEqual(headings, ["Experience", "Skills"])


if __name__ == "__main__":
    unittest.main()
