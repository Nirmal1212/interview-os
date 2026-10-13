"""Reading and writing resumes through one adapter.

read_resume() returns a list of Block objects from a .docx or PDF. The
workflow never touches python-docx or pypdf directly.

write_resume_docx() builds a single-column, ATS-safe .docx: no tables, text
boxes, or images, and headings use real heading styles.
"""

import os
import re
from dataclasses import dataclass

from docx import Document
from docx.shared import Pt
from pypdf import PdfReader

# Below this many extracted characters, a PDF is treated as having no text
# layer (for example a scanned image) and is rejected as a blocking issue.
MIN_TEXT_CHARS = 50

BLOCKING_MESSAGE = (
    "This resume has no readable text (it may be a scanned image or a PDF "
    "without a text layer). Re-export it as a .docx or a text-based PDF and try again."
)

KNOWN_HEADINGS = {
    "summary",
    "profile",
    "experience",
    "work experience",
    "professional experience",
    "skills",
    "technical skills",
    "education",
    "projects",
    "certifications",
    "achievements",
}


class UnreadableResumeError(ValueError):
    """Raised when a resume has no usable text. Carries the blocking-issue message."""


@dataclass(frozen=True)
class Block:
    kind: str  # "heading", "bullet", or "paragraph"
    text: str


def read_resume(path):
    """Read a .docx or PDF resume into Blocks."""
    extension = os.path.splitext(path)[1].lower()
    if extension == ".docx":
        blocks = _read_docx(path)
    elif extension == ".pdf":
        blocks = _read_pdf(path)
    else:
        raise ValueError("resume must be a .docx or .pdf file")
    total_chars = sum(len(b.text) for b in blocks)
    if total_chars < MIN_TEXT_CHARS:
        raise UnreadableResumeError(BLOCKING_MESSAGE)
    return blocks


def _read_docx(path):
    document = Document(path)
    blocks = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style is not None else ""
        if style.startswith("Heading") or style == "Title":
            blocks.append(Block("heading", text))
        elif style.startswith("List"):
            blocks.append(Block("bullet", text))
        else:
            blocks.append(Block(_classify_line(text), _strip_bullet(text)))
    return blocks


def _read_pdf(path):
    reader = PdfReader(path)
    blocks = []
    for page in reader.pages:
        for line in (page.extract_text() or "").splitlines():
            text = line.strip()
            if not text:
                continue
            blocks.append(Block(_classify_line(text), _strip_bullet(text)))
    return blocks


def _classify_line(text):
    if text.lower().rstrip(":") in KNOWN_HEADINGS:
        return "heading"
    if text.startswith(("- ", "• ", "* ")):
        return "bullet"
    return "paragraph"


def _strip_bullet(text):
    for marker in ("- ", "• ", "* "):
        if text.startswith(marker):
            return text[len(marker):]
    return text


def write_resume_docx(path, blocks):
    """Write Blocks as a single-column, ATS-safe .docx at path."""
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for block in blocks:
        if block.kind == "heading":
            document.add_heading(block.text, level=2)
        elif block.kind == "bullet":
            document.add_paragraph(block.text, style="List Bullet")
        else:
            document.add_paragraph(block.text)
    document.save(path)


_UNSAFE_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|]')


def ats_filename(candidate_name, role):
    """Return <CandidateName>_<Role>_ATS.docx with spaces replaced by underscores."""
    parts = []
    for part in (candidate_name, role):
        cleaned = _UNSAFE_FILENAME_CHARS.sub("", part).strip()
        parts.append("_".join(cleaned.split()))
    if not all(parts):
        raise ValueError("candidate name and role are both required for the file name")
    return f"{parts[0]}_{parts[1]}_ATS.docx"


def write_tailored_resume(master_path, out_dir, candidate_name, role, blocks):
    """Write the tailored resume into out_dir and return its path.

    Refuses to write if the output would land on the master file, so the
    master is never modified.
    """
    output_path = os.path.join(out_dir, ats_filename(candidate_name, role))
    if os.path.exists(output_path) and os.path.samefile(output_path, master_path):
        raise ValueError("output path is the master resume; refusing to overwrite it")
    if os.path.abspath(output_path) == os.path.abspath(master_path):
        raise ValueError("output path is the master resume; refusing to overwrite it")
    write_resume_docx(output_path, blocks)
    return output_path
